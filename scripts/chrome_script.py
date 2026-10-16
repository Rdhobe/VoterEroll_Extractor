"""Chrome Screen AI roll pipeline: PDF -> Markdown -> Excel.

Stage 1 (extraction): local OCR via Chrome's built-in Screen AI engine
(github.com/ayismas/chrome-ocr) - no API key, no cloud. Windows + Google
Chrome required.

Stage 2 (structuring): parses the Markdown into a voter Excel sheet
(merged from scripts/roll_to_excel.py, now kept as reference only).

Page convention for these rolls: page 1 = summary, page 2 = images only
(skipped), data starts at page 3. Override with --pages / --skip-pages.

Usage:
    python scripts/chrome_script.py --input "D:\\...\\roll.pdf" --output out.md
    python scripts/chrome_script.py --input "D:\\...\\roll.pdf" --output out.md --pages 3-10 --dpi 300 --excel voters.xlsx
"""
import argparse
import contextlib
import logging
import os
import re
import sys
import time
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import pymupdf as fitz
except ImportError:  # fallback keeps old environments working
    import fitz


def _silence_fitz_shim():
    """Swallow pymupdf's one-time `import fitz` deprecation notice, safely.

    Background: chrome_ocr lazily does `import fitz` inside ocr_pdf; the shim
    prints via pymupdf's internal print(flush=1), which crashed with
    OSError (WinError 1) when stdout was OS-redirected. Fix: point pymupdf's
    official message sink at devnull (set_messages API), then pre-import fitz
    so the in-library import later hits sys.modules cache. Real errors still
    raise as exceptions; only notice-level messages are dropped.
    """
    try:
        import pymupdf
        if hasattr(pymupdf, "set_messages"):
            pymupdf.set_messages(stream=open(os.devnull, "w"))
    except ImportError:
        pass
    if "fitz" in sys.modules:
        return
    try:
        import fitz  # noqa: F401
    except ImportError:
        pass


@contextlib.contextmanager
def quiet_stderr(quiet=True):
    """Redirect the Chrome DLL's native (C-level) stderr spam to devnull.

    GLOG env vars don't apply - the DLL logs straight to fd 2. Stdout (fd 1)
    is deliberately left alone: redirecting it breaks pymupdf's internal
    print() on Windows (OSError WinError 1, kills the run).
    --verbose keeps everything.
    """
    if not quiet:
        yield
        return
    devnull = os.open(os.devnull, os.O_WRONLY)
    old = os.dup(2)
    os.dup2(devnull, 2)
    try:
        yield
    finally:
        os.dup2(old, 2)
        os.close(devnull)
        os.close(old)


LOG_FMT = "%(asctime)s | %(levelname)-7s | %(message)s"


def setup_logger(log_file=None, verbose=False):
    logger = logging.getLogger("chromeocr")
    logger.setLevel(logging.DEBUG if verbose else logging.INFO)
    logger.handlers.clear()
    ch = logging.StreamHandler(sys.stdout)
    ch.setFormatter(logging.Formatter(LOG_FMT))
    logger.addHandler(ch)
    if log_file:
        Path(log_file).parent.mkdir(parents=True, exist_ok=True)
        fh = logging.FileHandler(log_file, encoding="utf-8")
        fh.setFormatter(logging.Formatter(LOG_FMT))
        logger.addHandler(fh)
    return logger


def parse_pages(spec, total=None):
    """'14,15' / '14-16' / 'all' -> sorted list of 1-based page numbers."""
    if not spec or spec.strip().lower() == "all":
        return None
    pages = set()
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            pages.update(range(int(a), int(b) + 1))
        elif part:
            pages.add(int(part))
    pages = sorted(pages)
    if total:
        pages = [p for p in pages if 1 <= p <= total]
    return pages


# ------------------------------------------------- stage 2: MD -> Excel
# Merged from scripts/roll_to_excel.py. The OCR dumps each page as separate,
# shuffled streams (serials / EPICs / card bodies) that are NOT next to each
# other, so re-linking is best-effort and every inferred link is flagged in
# the "Check" column for manual review.
DEV_DIGITS = str.maketrans("०१२३४५६७८९", "0123456789")

RE_SERIAL = re.compile(r"^\s*(\d{1,4})\s*$", re.M)          # a line holding only a number
RE_EPIC = re.compile(r"^\s*([A-Z]{2,4}\d{6,8}|UP/\d+/\d+/\d+)\s*$", re.M)

# card body fields (the 'ः' visarga is often OCR'd instead of ':')
COLON = r"\s*[:ः]+\s*"
RE_NAME = re.compile(rf"^\s*नाम{COLON}(.+?)\s*$", re.M)
RE_REL = re.compile(rf"^\s*(पिता|पति|माता|अन्य)\s*(?:का\s*)+नाम{COLON}(.+?)\s*$", re.M)
RE_HOUSE = re.compile(rf"मकान\s*संख्या{COLON}(\S+)")
RE_AGE_SEX = re.compile(rf"आयु{COLON}(\d+)\s*लिंग{COLON}(पुरुष|महिला|अन्य|\S+)")

# page header fields
RE_ASSEMBLY = re.compile(rf"विधानसभा\s*निर्वाचन\s*क्षेत्र[ःः:]*\s*(.+)")
RE_SECTION = re.compile(rf"अनुभाग\s*संख्या\s*और\s*नाम{COLON}(.+)")
RE_PART = re.compile(rf"भाग\s*संख्या{COLON}(\d+)")


def clean_md(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip(" ):(")


def parse_md_page(page_no: int, txt: str) -> list:
    txt = txt.translate(DEV_DIGITS)  # normalize Devanagari digits up front
    txt = re.sub(r"^\s*#+\s*", "", txt, flags=re.M)  # strip '### ' heading marks
    serials = RE_SERIAL.findall(txt)

    # ---- 1. card bodies: cut the page at every Name line, parse WITHIN ----
    # Index-zipping parallel field lists breaks the moment one list gains or
    # loses a match mid-page (every card after the gap gets a neighbour's
    # relation/house/sex). Chunking keeps each card's fields together; a
    # missing field only empties that card's slot (+flag), never shifts others.
    name_hits = list(RE_NAME.finditer(txt))
    if not name_hits:
        if txt.strip():
            print(f"  ! page {page_no}: no Name lines found", file=sys.stderr)
        return []
    cards = []
    for k, m in enumerate(name_hits):
        chunk = txt[m.start():name_hits[k + 1].start() if k + 1 < len(name_hits) else len(txt)]
        name = clean_md(m.group(1))
        rm = RE_REL.search(chunk)
        hm = RE_HOUSE.search(chunk)
        am = RE_AGE_SEX.search(chunk)
        notes = []
        rel_type, rel_name = (rm.group(1), re.sub(r"^नाम\s*[:ः]\s*", "", clean_md(rm.group(2)))) if rm else ("", "")
        if rm and RE_REL.search(chunk, rm.end()):
            notes.append("2nd relation in block ignored")
        house = clean_md(hm.group(1)) if hm else ""
        if not hm:
            notes.append("house missing")
        if am:
            age, gender = int(am.group(1)), am.group(2)
        else:
            age, gender = "", ""
            notes.append("age/sex missing")
        if rel_name and rel_name == name:
            # a relative is never the voter: neighbour fragment, drop it
            rel_name = ""
            notes.append("relation==name dropped")
        if rel_type == "पति" and gender == "पुरुष":
            notes.append("relation/gender conflict")
        cards.append(dict(Name=name, Relation_Type=rel_type, Relation_Name=rel_name,
                          House=house, Age=age, Gender=gender, _notes=notes,
                          _file_pos=k))
    n = len(cards)

    # ---- 2. serials: sorted, they are the consecutive block on the page ----
    # Keep the largest gap-tolerant cluster (gap <= 5): strays like the
    # भाग-संख्या "1" (OCR splits "भाग संख्या :: 1" across lines) form their own
    # lone cluster and are dropped, while a genuine single gap (e.g. a deleted
    # serial mid-page) survives. Clean pages are one cluster: unchanged behavior.
    serials = sorted(set(int(s) for s in serials))
    clusters, cur = [], []
    for s in serials:
        if cur and s - cur[-1] > 5:
            clusters.append(cur)
            cur = []
        cur.append(s)
    if cur:
        clusters.append(cur)
    best = max(clusters, key=len) if clusters else []
    stray = sorted(set(serials) - set(best))
    if stray:
        print(f"  ! page {page_no}: ignoring stray serials {stray}", file=sys.stderr)
    serials = best
    if len(serials) != n:
        print(f"  ! page {page_no}: {len(serials)} serials vs {n} cards", file=sys.stderr)

    # House numbers never decrease with serial number, but the OCR shuffles
    # card order -> stable-sort the cards by house number, then hand out serials.
    def house_key(c):
        m = re.match(r"\d+", c["House"])
        return int(m.group()) if m else 10**9
    ordered = sorted(cards, key=house_key)
    group_size = defaultdict(int)
    for c in cards:
        group_size[house_key(c)] += 1
    for pos, c in enumerate(ordered):
        if pos < len(serials):
            c["Serial"] = serials[pos]
            if c["_file_pos"] != pos:
                c["_serial_note"] = "serial inferred"
            elif group_size[house_key(c)] > 1:
                # block of serials for this house is right, order inside it is OCR order
                c["_serial_note"] = "verify order within house"
            else:
                c["_serial_note"] = ""
        else:  # more cards than serial lines on this page - keep the card, flag it
            c["Serial"] = None
            c["_serial_note"] = "serial missing"

    # ---- 3. EPIC IDs: pair a serial with the token that directly follows it
    stream = [(m.start(), "S", m.group(1)) for m in RE_SERIAL.finditer(txt)] + \
             [(m.start(), "E", m.group(1)) for m in RE_EPIC.finditer(txt)]
    stream.sort()
    epic_of, used = {}, set()
    for (_, k1, v1), (_, k2, v2) in zip(stream, stream[1:]):
        if k1 == "S" and k2 == "E":
            epic_of[int(v1)] = v2
            used.add(v2)
    leftovers = [v for _, k, v in stream if k == "E" and v not in used]
    for s in serials:                                # fill gaps in stream order
        if s not in epic_of and leftovers:
            epic_of[s] = leftovers.pop(0)
            epic_of[("guess", s)] = True
    for c in cards:
        s = c.get("Serial")
        c["EPIC"] = epic_of.get(s, "") if s is not None else ""
        c["_epic_note"] = "EPIC inferred" if (s is not None and epic_of.get(("guess", s))) else ""

    # ---- 4. page header info ---------------------------------------------
    asm = RE_ASSEMBLY.search(txt); sec = RE_SECTION.search(txt); prt = RE_PART.search(txt)
    for c in cards:
        c["Page"] = page_no
        c["Assembly"] = clean_md(asm.group(1)) if asm else ""
        c["Section"] = clean_md(sec.group(1)) if sec else ""
        c["Part"] = prt.group(1) if prt else ""
        c["Check"] = "; ".join(x for x in [c.pop("_serial_note"), c.pop("_epic_note")] + c.pop("_notes") if x)
        c.pop("_file_pos")
    return cards


def write_blocked_md(per_page_info, md_path, stem, dpi, pages_spec, dt):
    """Write one voter block per card (photo-caption noise dropped).

    per_page_info = [(page_no, assembly, section, part, cards)].
    """
    out = Path(md_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    L = []
    L.append("# Source: %s" % stem)
    L.append("# Engine: Chrome Screen AI | dpi=%d | pages=%s | %.1fs" % (dpi, pages_spec, dt))
    L.append("")
    for pno, asm, sec, prt, cards in per_page_info:
        L.append("=" * 50)
        L.append("--- PAGE %d ---" % pno)
        meta = " | ".join(x for x in (asm, sec, ("Part " + prt) if prt else "") if x)
        if meta:
            L.append("# " + meta)
        L.append("=" * 50)
        L.append("")
        if not cards:
            L.append("(no voter rows on this page)")
            L.append("")
            continue
        for i, c in enumerate(cards, 1):
            L.append("--- BLOCK %d ---" % i)
            L.append("Serial: %s | EPIC: %s" % (c.get("Serial") or "", c.get("EPIC") or ""))
            L.append("नाम : %s" % (c.get("Name") or ""))
            L.append("%s का नाम : %s" % (c.get("Relation_Type") or "", c.get("Relation_Name") or ""))
            L.append("मकान संख्या : %s" % (c.get("House") or ""))
            L.append("आयु : %s | लिंग : %s" % (c.get("Age") or "", c.get("Gender") or ""))
            if c.get("Check"):
                L.append("Check: %s" % c["Check"])
            L.append("------------------------------")
            L.append("")
    out.write_text("\n".join(L), encoding="utf-8")


def cards_to_excel(all_cards, xlsx_path, logger):
    """Formatted voter workbook straight from parsed cards (no re-parsing)."""
    if not all_cards:
        logger.warning("No voter rows to write.")
        return 0, 0
    cols = ["Page", "Part", "Section", "Serial", "EPIC", "Name", "Relation_Type",
            "Relation_Name", "House", "Age", "Gender", "Check", "Assembly"]
    df = pd.DataFrame(all_cards, columns=cols).sort_values(["Serial", "Page"], na_position="last")
    df.columns = ["Page", "Part No", "Section", "Serial No", "EPIC ID", "Name",
                  "Relation", "Relative Name", "House No", "Age", "Gender", "Check", "Assembly"]
    Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(xlsx_path, index=False, sheet_name="Voters")

    wb = load_workbook(xlsx_path); ws = wb["Voters"]
    hdr_fill = PatternFill("solid", start_color="1F4E78")
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")
    flag = PatternFill("solid", start_color="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
        if row[11].value:                                  # 'Check' column
            for c in row:
                c.fill = flag
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 40)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(xlsx_path)
    flagged = int((df["Check"] != "").sum())
    return len(df), flagged


def build_parser():
    p = argparse.ArgumentParser(description="Roll PDF -> Markdown (Screen AI) -> voter Excel.")
    p.add_argument("--input", required=True, help="Input PDF file")
    p.add_argument("--output", required=True, help="Output Markdown file")
    p.add_argument("--pages", default="all", help="'all' or e.g. '3-10' / '14,15' (1-based)")
    p.add_argument("--skip-pages", default="1,2",
                   help="Pages to skip: p1 = summary, p2 = images only (default '1,2', '' to disable)")
    p.add_argument("--excel", default=None,
                   help="Voter workbook path (default: <output-stem>.xlsx next to the MD)")
    p.add_argument("--no-excel", action="store_true", help="Skip stage 2 (Markdown only)")
    p.add_argument("--keep-raw-md", action="store_true",
                   help="Also save the shuffled per-page Markdown as <output-stem>.raw.md")
    p.add_argument("--dpi", type=int, default=200)
    p.add_argument("--log-file", default=None)
    p.add_argument("--verbose", action="store_true")
    return p


def extract_pdf_to_md(pdf_path, pages_spec="all", skip_pages="1,2",
                      dpi=200, raw_path=None, logger=None, verbose=False):
    """Stage 1: PDF -> per-page Markdown via Chrome Screen AI.

    Importable (test scripts reuse it without the CLI).
    Returns (paged, n_chars, seconds) with paged = [(page_no, body_md)].
    Writes the raw shuffled Markdown only if raw_path is given.
    Raises FileNotFoundError / RuntimeError.
    """
    log = logger or logging.getLogger("chromeocr")
    _silence_fitz_shim()  # idempotent; also covers non-CLI importers
    try:
        with quiet_stderr(not verbose):
            from chrome_ocr import ScreenAIEngine, ocr_pdf
    except ImportError as e:
        raise RuntimeError("chrome-ocr not installed. Run: pip install -r requirements.txt") from e

    pdf = Path(pdf_path)
    if not pdf.is_file():
        raise FileNotFoundError(str(pdf))

    log.info("Loading Screen AI engine...")
    t0 = time.time()
    try:
        with quiet_stderr(not verbose):
            engine = ScreenAIEngine()
    except Exception as e:
        raise RuntimeError("Engine init failed: %s" % e) from e
    log.info("Engine ready in %.1fs", time.time() - t0)

    pages = parse_pages(pages_spec)
    log.info("Input=%s pages=%s dpi=%d", pdf.name, pages_spec, dpi)

    if pages is None:  # expand "all" so every page gets its own tagged section
        with fitz.open(str(pdf)) as doc:
            pages = list(range(1, doc.page_count + 1))

    skip = parse_pages(skip_pages) or []  # p1 = summary, p2 = images only
    if skip:
        pages = [pn for pn in pages if pn not in set(skip)]
        log.info("Skipping pages %s (summary/image-only)", skip_pages)
    if not pages:
        raise RuntimeError("Nothing left after skip-pages=%s" % skip_pages)

    t1 = time.time()
    paged, n_chars = [], 0
    try:
        for i, pno in enumerate(pages, 1):
            with quiet_stderr(not verbose):
                md = ocr_pdf(str(pdf), dpi=dpi, pages=[pno], engine=engine)
            paged.append((pno, md or ""))
            n_chars += len(md or "")
            if i % 5 == 0 or i == len(pages):
                log.info("%d/%d pages OCR'd", i, len(pages))
    except Exception:
        log.exception("Extraction failed")
        raise
    dt = time.time() - t1

    if raw_path:  # optional audit copy of the shuffled per-page Markdown
        out = Path(raw_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        header = ("# Source: %s\n# Engine: Chrome Screen AI | dpi=%d | pages=%s | %.1fs\n\n"
                  % (pdf.name, dpi, pages_spec, dt))
        out.write_text(header + "\n\n".join("--- PAGE %d ---\n\n%s" % (p, b)
                                            for p, b in paged), encoding="utf-8")

    log.info("DONE: %d pages in %.1fs (%.1fs/page), %d chars",
             len(pages), dt, dt / len(pages), n_chars)
    if not n_chars:
        log.warning("Empty output - page range may exceed the PDF or DLL failed.")
    return paged, n_chars, dt


def main(argv=None):
    args = build_parser().parse_args(argv)
    logger = setup_logger(args.log_file, args.verbose)

    _silence_fitz_shim()  # pre-fire chrome_ocr's lazy `import fitz` (see helper)

    pdf = Path(args.input)
    out = Path(args.output)
    try:
        paged, n_chars, dt = extract_pdf_to_md(
            str(pdf), args.pages, args.skip_pages, args.dpi,
            raw_path=(str(out.with_suffix(".raw.md")) if args.keep_raw_md else None),
            logger=logger, verbose=args.verbose)
    except FileNotFoundError:
        logger.error("Not found: %s", pdf)
        return 1
    except RuntimeError as e:
        logger.error("%s", e)
        return 1
    if not n_chars:
        return 1

    # ---- stage 2: raw sections -> cards -> BLOCKED md -> Excel --------------
    per_page, all_cards = [], []
    for pno, body in paged:
        cards = parse_md_page(pno, body)
        all_cards.extend(cards)
        if cards:
            c0 = cards[0]
            per_page.append((pno, c0.get("Assembly") or "", c0.get("Section") or "",
                             c0.get("Part") or "", cards))
        else:
            per_page.append((pno, "", "", "", []))
    write_blocked_md(per_page, str(out), pdf.stem, args.dpi, args.pages, dt)
    logger.info("BLOCKS: %d voter blocks -> %s", len(all_cards), out.name)

    if args.no_excel:
        return 0
    xlsx = args.excel or str(out.with_suffix(".xlsx"))
    try:
        n_rows, n_flag = cards_to_excel(all_cards, xlsx, logger)
    except PermissionError:
        logger.error("CLOSE %s in Excel and re-run (file locked).", Path(xlsx).name)
        return 1
    logger.info("TABLE: %d voter rows (%d flagged for review) -> %s", n_rows, n_flag,
                Path(xlsx).name)
    return 0


if __name__ == "__main__":
    sys.exit(main())
